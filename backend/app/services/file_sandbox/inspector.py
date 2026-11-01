"""Structural inspection of one file, run inside the file sandbox.

Importable as ``heym_inspector`` in the sandbox, where pandas, openpyxl, xlrd,
pdfplumber, pypdf, python-docx and langdetect are installed, and ``heym_ocr`` (the
Converter node's Tesseract OCR) reads scanned pages. It reads only what a summary
needs: the first pages, ``read_only`` workbooks, a sample of rows, the first
records of a large JSON file. Never imported by the backend at run time.
"""

from __future__ import annotations

import csv
import datetime as _dt
import json
import re
from pathlib import Path
from typing import Any

SAMPLE_CHARS = 2000
SAMPLE_ROWS = 5
DTYPE_ROWS = 200
CSV_STATS_ROWS = 10_000
JSON_FULL_LOAD_BYTES = 50 * 1024 * 1024
MAX_PAGES_PER_CALL = 10
MAX_OCR_PAGES = 5
MAX_OUTLINE = 60
HEADER_SCAN_ROWS = 30
MAX_TITLE_ROWS = 5
MAX_KEY_VALUES = 20

_EXTENSION_KINDS = {
    ".pdf": "pdf",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".xls": "xls",
    ".csv": "csv",
    ".tsv": "csv",
    ".docx": "docx",
    ".json": "json",
    ".jsonl": "jsonl",
    ".ndjson": "jsonl",
    ".txt": "text",
    ".md": "text",
    ".log": "text",
    ".xml": "text",
    ".html": "text",
    ".htm": "text",
}
_UTF8_BOM = b"\xef\xbb\xbf"


def _clip(text: str | None, limit: int = SAMPLE_CHARS) -> str:
    text = text or ""
    return text if len(text) <= limit else text[:limit] + " ..."


def _jsonable(value: Any) -> Any:
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    if isinstance(value, float) and value != value:
        return None
    if isinstance(value, (str, int, float, bool)) or value is None:
        return _clip(value, 300) if isinstance(value, str) else value
    return _clip(str(value), 300)


def _language(text: str) -> dict[str, Any] | None:
    sample = re.sub(r"\s+", " ", text or "").strip()[:5000]
    if len(sample) < 20:
        return None
    try:
        from langdetect import DetectorFactory, detect_langs

        DetectorFactory.seed = 0
        best = detect_langs(sample)[0]
    except Exception:
        return None
    return {"code": best.lang, "confidence": round(float(best.prob), 2)}


def detect_kind(path: Path, name: str | None = None, mime: str | None = None) -> str:
    """File kind from the extension first, then from the first bytes."""
    suffix = Path(name or path.name).suffix.lower()
    if suffix in _EXTENSION_KINDS:
        return _EXTENSION_KINDS[suffix]
    with open(path, "rb") as handle:
        head = handle.read(4096)
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    if head.startswith(b"PK"):
        import zipfile

        try:
            names = zipfile.ZipFile(path).namelist()
        except zipfile.BadZipFile:
            return "binary"
        if any(n.startswith("xl/") for n in names):
            return "xlsx"
        if any(n.startswith("word/") for n in names):
            return "docx"
        return "zip"
    stripped = head.lstrip()
    if stripped[:1] in (b"{", b"["):
        return "json"
    if mime and mime.startswith("text/"):
        return "text"
    try:
        head.decode("utf-8")
        return "text"
    except UnicodeDecodeError:
        return "binary"


def _parse_pages(spec: Any, page_count: int) -> list[int]:
    """``"1-3,5"`` or ``[1, 2]`` to 1-based page numbers, capped per call."""
    pages: list[int] = []
    if isinstance(spec, list):
        parts = [str(p) for p in spec]
    else:
        parts = [p for p in re.split(r"[,\s]+", str(spec or "")) if p]
    for part in parts:
        if "-" in part:
            start_s, _, end_s = part.partition("-")
            try:
                start, end = int(start_s), int(end_s or page_count)
            except ValueError:
                continue
            pages.extend(range(start, end + 1))
        else:
            try:
                pages.append(int(part))
            except ValueError:
                continue
    unique = [p for p in dict.fromkeys(pages) if 1 <= p <= page_count]
    return unique[:MAX_PAGES_PER_CALL]


def _inspect_pdf(path: Path, options: dict[str, Any]) -> dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    encrypted = bool(reader.is_encrypted)
    if encrypted:
        try:
            opened = reader.decrypt("")
        except Exception:
            opened = 0
        if not opened:
            return {
                "encrypted": True,
                "error": "The PDF is password protected.",
                "suggestion": "Ask for an unprotected copy, or the password to remove it first.",
            }
    page_count = len(reader.pages)
    meta = reader.metadata or {}
    metadata = {
        key: _jsonable(meta.get(f"/{key.capitalize()}"))
        for key in ("title", "author", "creator", "producer", "subject")
        if meta.get(f"/{key.capitalize()}")
    }
    sizes: list[list[float]] = []
    for page in reader.pages[:20]:
        box = page.mediabox
        size = [round(float(box.width), 1), round(float(box.height), 1)]
        if size not in sizes:
            sizes.append(size)

    import pdfplumber

    requested = _parse_pages(options.get("pages"), page_count) if options.get("pages") else []
    probe = sorted(set(requested or [1, 2]) & set(range(1, page_count + 1)))
    pages_out: list[dict[str, Any]] = []
    probe_text = ""
    with pdfplumber.open(str(path)) as pdf:
        for number in probe:
            page = pdf.pages[number - 1]
            text = page.extract_text() or ""
            probe_text += text
            if requested or number == 1:
                tables = page.extract_tables() or []
                entry: dict[str, Any] = {
                    "page": number,
                    "chars": len(text),
                    "text": _clip(text),
                    "tables": len(tables),
                }
                if tables:
                    entry["first_table_rows"] = [
                        [_jsonable(cell) for cell in row] for row in tables[0][:SAMPLE_ROWS]
                    ]
                pages_out.append(entry)
    extractable = len(probe_text.strip()) >= 20
    result: dict[str, Any] = {
        "page_count": page_count,
        "encrypted": encrypted,
        "metadata": metadata,
        "page_sizes_pt": sizes,
        "text_extractable": extractable,
        "language": _language(probe_text),
        "pages": pages_out,
    }
    if not extractable:
        result["scanned"] = True
        result.update(_ocr_pages(path, requested or probe))
    return result


def _ocr_engine() -> Any:
    """The Converter node's OCR: ``heym_ocr`` in the sandbox, the service in tests."""
    try:
        import heym_ocr  # type: ignore[import-not-found]
    except ImportError:
        from app.services import ocr_service as heym_ocr
    return heym_ocr


def _page_runs(pages: list[int]) -> list[tuple[int, int]]:
    runs: list[tuple[int, int]] = []
    for number in sorted(pages):
        if runs and number == runs[-1][1] + 1:
            runs[-1] = (runs[-1][0], number)
        else:
            runs.append((number, number))
    return runs


def _ocr_pages(path: Path, pages: list[int]) -> dict[str, Any]:
    """Read scanned pages with Tesseract, at most ``MAX_OCR_PAGES`` per call."""
    engine = _ocr_engine()
    data = path.read_bytes()
    language = engine.AUTO_LANGUAGE
    read: list[dict[str, Any]] = []
    try:
        for first, last in _page_runs(pages[:MAX_OCR_PAGES]):
            ocr = engine.pdf_to_text(data, language=language, page_range=f"{first}-{last}")
            language = ocr.language
            read.extend(
                {"page": page.page, "chars": len(page.text), "text": _clip(page.text), "ocr": True}
                for page in ocr.pages
            )
    except ValueError as exc:
        return {
            "ocr_error": str(exc),
            "suggestion": "No text layer and OCR failed here. The Converter node's pdfToText "
            "reads scanned PDFs in a workflow.",
        }
    return {
        "pages": read,
        "ocr_language": language,
        "language": _language(" ".join(page["text"] for page in read)),
        "suggestion": f"No text layer: these pages were read with OCR. Pass `pages` for others "
        f"(at most {MAX_OCR_PAGES} per call). In a workflow, the Converter node's pdfToText "
        "does the same.",
    }


def _infer_type(values: list[Any]) -> str:
    seen = {
        "empty"
        if v is None or v == ""
        else "bool"
        if isinstance(v, bool)
        else "int"
        if isinstance(v, int)
        else "float"
        if isinstance(v, float)
        else "datetime"
        if isinstance(v, (_dt.datetime, _dt.date))
        else "string"
        for v in values
    }
    seen.discard("empty")
    if not seen:
        return "empty"
    if seen == {"int", "float"}:
        return "float"
    return seen.pop() if len(seen) == 1 else "mixed"


def _filled(row: tuple[Any, ...]) -> list[tuple[int, Any]]:
    return [
        (index, value)
        for index, value in enumerate(row)
        if value is not None and not (isinstance(value, str) and not value.strip())
    ]


def _label_like(value: Any) -> bool:
    """Text, a date, or a year: what a column header holds."""
    if isinstance(value, (str, _dt.datetime, _dt.date)):
        return True
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return False
    return float(value).is_integer() and 1900 <= value <= 2100


def _value_kind(value: Any) -> str:
    if isinstance(value, (_dt.datetime, _dt.date)):
        return "date"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "number"
    return type(value).__name__


def _find_header(rows: list[tuple[Any, ...]]) -> int | None:
    """Index of the header row, skipping title rows, notes and blank rows above it.

    A header holds at least two mostly-text cells and spans the table below it. With
    two cells, the second column needs at least two values below it, all of one kind,
    so a summary sheet of label/value pairs is not mistaken for a two-column table.
    """
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        cells = _filled(row)
        if len(cells) < 2 or sum(_label_like(v) for _i, v in cells) < 0.8 * len(cells):
            continue
        below = [found for found in (_filled(r) for r in rows[index + 1 : index + 51]) if found]
        below = below[:5]
        if not below or len(cells) < 0.6 * max(len(found) for found in below):
            continue
        has_values = any(not isinstance(v, str) for found in below for _i, v in found)
        if not has_values and len(cells) < 3:
            continue
        if len(cells) == 2:
            second = cells[1][0]
            values = [v for found in below for i, v in found if i == second]
            kinds = {_value_kind(v) for v in values}
            if len(values) < 2 or len(kinds) != 1 or kinds == {"str"}:
                continue
        return index
    return None


def _column_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, rest = divmod(index - 1, 26)
        letters = chr(65 + rest) + letters
    return letters


def _title_rows(rows: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    titles = []
    for index, row in enumerate(rows):
        cells = _filled(row)
        if cells:
            text = " | ".join(str(_jsonable(v)) for _i, v in cells)
            titles.append({"row": index + 1, "text": text})
    return titles[:MAX_TITLE_ROWS]


def _key_values(rows: list[tuple[Any, ...]]) -> tuple[list[dict[str, Any]], int] | None:
    """Label/value pairs of a summary sheet ("Total | 1266"), or None for other layouts."""
    filled_rows = [(index, _filled(row)) for index, row in enumerate(rows)]
    filled_rows = [(index, cells) for index, cells in filled_rows if cells]
    pairs = []
    first_pair = None
    for index, cells in filled_rows:
        if len(cells) == 2 and isinstance(cells[0][1], str):
            first_pair = index if first_pair is None else first_pair
            (_label_col, label), (value_col, value) = cells
            pairs.append(
                {
                    "cell": f"{_column_letter(value_col)}{index + 1}",
                    "label": _jsonable(label),
                    "value": _jsonable(value),
                }
            )
    if first_pair is None or len(pairs) < 2:
        return None
    rest = [cells for index, cells in filled_rows if index >= first_pair]
    if len(pairs) < 0.6 * len(rest):
        return None
    return pairs[:MAX_KEY_VALUES], first_pair


def _table_summary(rows: list[tuple[Any, ...]], sheet_name: str | None = None) -> dict[str, Any]:
    """Title rows, header, column types and sample rows from the first rows of a sheet."""
    header_index = _find_header(rows)
    if header_index is None:
        layout = _key_values(rows)
        if layout is not None:
            pairs, first_pair = layout
            return {
                "layout": "key_value",
                "header_row": None,
                "title_rows": _title_rows(rows[:first_pair]),
                "key_values": pairs,
            }
        header_index = next((i for i, row in enumerate(rows[:10]) if _filled(row)), None)
    if header_index is None:
        return {"header_row": None, "columns": [], "sample_rows": []}
    header = rows[header_index]
    header_cells = _filled(header)
    first_col = min(i for i, _v in header_cells)
    last_col = max(
        [i for i, _v in header_cells]
        + [i for row in rows[header_index + 1 :] for i, _v in _filled(row)]
    )
    body = [row for row in rows[header_index + 1 :] if _filled(row)]
    columns = []
    for index in range(first_col, last_col + 1):
        title = header[index] if index < len(header) else None
        values = [row[index] if index < len(row) else None for row in body[:DTYPE_ROWS]]
        columns.append(
            {
                "name": _jsonable(title) if title not in (None, "") else f"column_{index + 1}",
                "type": _infer_type(values),
            }
        )
    summary: dict[str, Any] = {
        "header_row": header_index + 1,
        "title_rows": _title_rows(rows[:header_index]),
        "columns": columns,
        "sample_rows": [
            [_jsonable(row[i] if i < len(row) else None) for i in range(first_col, last_col + 1)]
            for row in body[:SAMPLE_ROWS]
        ],
    }
    if first_col:
        summary["first_column"] = _column_letter(first_col)
    if sheet_name is not None:
        hint = f"pd.read_excel(path, sheet_name={sheet_name!r}, header={header_index}"
        if first_col:
            hint += f', usecols="{_column_letter(first_col)}:{_column_letter(last_col)}"'
        summary["read_hint"] = hint + ")"
    return summary


def _inspect_xlsx(path: Path, options: dict[str, Any]) -> dict[str, Any]:
    from openpyxl import load_workbook

    wanted = options.get("sheet")
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        names = list(workbook.sheetnames)
        defined = []
        try:
            for name, item in workbook.defined_names.items():
                defined.append({"name": name, "refers_to": _jsonable(item.attr_text)})
        except Exception:
            defined = []
        scan_all = path.stat().st_size <= 5 * 1024 * 1024
        sheets = []
        for name in names:
            sheet = workbook[name]
            info: dict[str, Any] = {
                "name": name,
                "row_count": sheet.max_row,
                "column_count": sheet.max_column,
            }
            if wanted not in (None, "") and str(wanted) != name:
                sheets.append(info)
                continue
            first_rows = list(sheet.iter_rows(min_row=1, max_row=DTYPE_ROWS + 10, values_only=True))
            info.update(_table_summary(first_rows, name))
            if scan_all or str(wanted or "") == name:
                last_filled = 0
                for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if any(v not in (None, "") for v in row):
                        last_filled = index
                info["last_filled_row"] = last_filled
            sheets.append(info)
        if wanted not in (None, "") and str(wanted) not in names:
            return {
                "sheet_names": names,
                "error": f"No sheet named {wanted!r}.",
                "suggestion": f"Use one of: {', '.join(names)}",
            }
    finally:
        workbook.close()
    return {"sheet_names": names, "named_ranges": defined, "sheets": sheets}


def _inspect_xls(path: Path, options: dict[str, Any]) -> dict[str, Any]:
    import xlrd

    wanted = options.get("sheet")
    book = xlrd.open_workbook(str(path), on_demand=True)
    try:
        names = book.sheet_names()
        sheets = []
        for name in names:
            sheet = book.sheet_by_name(name)
            info: dict[str, Any] = {
                "name": name,
                "row_count": sheet.nrows,
                "column_count": sheet.ncols,
            }
            if wanted in (None, "") or str(wanted) == name:
                rows = [
                    tuple(sheet.row_values(i)) for i in range(min(sheet.nrows, DTYPE_ROWS + 10))
                ]
                info.update(_table_summary(rows, name))
            sheets.append(info)
            book.unload_sheet(name)
    finally:
        book.release_resources()
    return {"sheet_names": names, "sheets": sheets}


def _read_head(path: Path, size: int = 256 * 1024) -> bytes:
    with open(path, "rb") as handle:
        return handle.read(size)


def _detect_encoding(raw: bytes) -> str:
    """UTF-8 (with or without BOM), else cp1254, which keeps Turkish letters, else latin-1."""
    if raw.startswith(_UTF8_BOM):
        return "utf-8-sig"
    # A multi-byte character may be cut at the end of the sample.
    for trim in range(4):
        try:
            raw[: len(raw) - trim].decode("utf-8")
            return "utf-8"
        except UnicodeDecodeError:
            continue
    try:
        raw.decode("cp1254")
        return "cp1254"
    except UnicodeDecodeError:
        return "latin-1"


def _count_lines(path: Path) -> int:
    count = 0
    last = b""
    with open(path, "rb") as handle:
        while chunk := handle.read(1024 * 1024):
            count += chunk.count(b"\n")
            last = chunk
    if last and not last.endswith(b"\n"):
        count += 1
    return count


def _inspect_csv(path: Path, options: dict[str, Any]) -> dict[str, Any]:
    import pandas as pd

    raw = _read_head(path)
    encoding = _detect_encoding(raw)
    sample = raw.decode(encoding, errors="replace")
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        delimiter = csv.Sniffer().sniff(sample[:65536], delimiters=",;\t|").delimiter
    except csv.Error:
        pass
    frame = pd.read_csv(
        path, sep=delimiter, encoding=encoding, nrows=CSV_STATS_ROWS, low_memory=False
    )
    lines = _count_lines(path)
    exact = bool(options.get("full"))
    rows = lines - 1 if lines else 0
    nulls = frame.isna().mean()
    if exact:
        rows = 0
        null_counts = None
        for chunk in pd.read_csv(
            path, sep=delimiter, encoding=encoding, chunksize=100_000, low_memory=False
        ):
            rows += len(chunk)
            counts = chunk.isna().sum()
            null_counts = counts if null_counts is None else null_counts + counts
        if null_counts is not None and rows:
            nulls = null_counts / rows
    return {
        "encoding": encoding,
        "delimiter": delimiter,
        "rows": rows,
        "rows_exact": exact,
        "columns": [
            {
                "name": str(col),
                "dtype": str(frame[col].dtype),
                "null_ratio": round(float(nulls[col]), 4),
            }
            for col in frame.columns
        ],
        "null_ratio_basis": "all rows"
        if exact
        else f"first {min(len(frame), CSV_STATS_ROWS)} rows",
        "sample_rows": [
            [_jsonable(v) for v in row] for row in frame.head(SAMPLE_ROWS).itertuples(index=False)
        ],
    }


def _inspect_docx(path: Path, options: dict[str, Any]) -> dict[str, Any]:
    import docx

    document = docx.Document(str(path))
    paragraphs = document.paragraphs
    outline = []
    texts: list[str] = []
    for paragraph in paragraphs:
        text = paragraph.text or ""
        texts.append(text)
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if (style.startswith("Heading") or style == "Title") and text.strip():
            level_match = re.search(r"(\d+)$", style)
            level = int(level_match.group(1)) if level_match else 0
            if len(outline) < MAX_OUTLINE:
                outline.append({"level": level, "text": _clip(text, 200)})
    tables = []
    for table in document.tables:
        cells_text = [[cell.text for cell in row.cells] for row in table.rows]
        for row in cells_text:
            texts.extend(row)
        tables.append({"rows": len(table.rows), "columns": len(table.columns)})
    full_text = "\n".join(texts)
    props = document.core_properties
    first_table = (
        [
            [_clip(cell.text, 200) for cell in row.cells]
            for row in document.tables[0].rows[:SAMPLE_ROWS]
        ]
        if document.tables
        else []
    )
    return {
        "paragraphs": len(paragraphs),
        "headings": outline,
        "tables": tables,
        "first_table_rows": first_table,
        "words": len(full_text.split()),
        "characters": len(full_text),
        "metadata": {
            key: _jsonable(getattr(props, key, None))
            for key in ("title", "author", "last_modified_by", "created", "modified")
            if getattr(props, key, None)
        },
        "language": _language(full_text),
        "sample_text": _clip("\n".join(t for t in texts[:200] if t.strip())),
    }


def _depth(value: Any, level: int = 0) -> int:
    if level > 50:
        return level
    if isinstance(value, dict) and value:
        return max(_depth(v, level + 1) for v in list(value.values())[:200])
    if isinstance(value, list) and value:
        return max(_depth(v, level + 1) for v in value[:50])
    return level


def _type_name(value: Any) -> str:
    return {dict: "object", list: "array", str: "string", bool: "boolean", type(None): "null"}.get(
        type(value), "number" if isinstance(value, (int, float)) else type(value).__name__
    )


def _sample(value: Any) -> Any:
    text = json.dumps(value, default=str, ensure_ascii=False)
    return value if len(text) <= SAMPLE_CHARS else _clip(text)


def _describe_records(records: list[Any]) -> dict[str, Any]:
    keys: dict[str, str] = {}
    for item in records[:100]:
        if isinstance(item, dict):
            for key, val in item.items():
                keys.setdefault(str(key), _type_name(val))
    return {"record_keys": keys, "sample_records": [_sample(r) for r in records[:3]]}


def _stream_array_head(path: Path, encoding: str, count: int) -> list[Any]:
    decoder = json.JSONDecoder()
    items: list[Any] = []
    buffer = ""
    with open(path, encoding=encoding, errors="replace") as handle:
        buffer = handle.read(1024 * 1024).lstrip()
        if not buffer.startswith("["):
            return items
        position = 1
        while len(items) < count:
            while position < len(buffer) and buffer[position] in " \r\n\t,":
                position += 1
            if position < len(buffer) and buffer[position] == "]":
                break
            try:
                item, end = decoder.raw_decode(buffer, position)
            except json.JSONDecodeError:
                more = handle.read(1024 * 1024)
                if not more:
                    break
                buffer += more
                continue
            items.append(item)
            position = end
    return items


def _inspect_json(path: Path, options: dict[str, Any]) -> dict[str, Any]:
    encoding = _detect_encoding(_read_head(path, 64 * 1024))
    size = path.stat().st_size
    if size > JSON_FULL_LOAD_BYTES:
        head = _read_head(path, 4096).decode(encoding, errors="replace").lstrip()
        if head.startswith("["):
            items = _stream_array_head(path, encoding, 20)
            return {
                "root_type": "array",
                "length": None,
                "note": "Large file: only the first records were read.",
                "depth_estimate": _depth(items),
                **_describe_records(items),
            }
        return {
            "root_type": "object" if head.startswith("{") else "unknown",
            "note": "Large object: use run_python with a streaming read for its keys.",
        }
    with open(path, encoding=encoding) as handle:
        data = json.load(handle)
    result: dict[str, Any] = {"root_type": _type_name(data), "depth_estimate": _depth(data)}
    if isinstance(data, list):
        result["length"] = len(data)
        result.update(_describe_records(data))
    elif isinstance(data, dict):
        result["keys"] = {str(k): _type_name(v) for k, v in list(data.items())[:100]}
        result["key_count"] = len(data)
        arrays = {k: len(v) for k, v in data.items() if isinstance(v, list)}
        if arrays:
            result["array_lengths"] = dict(list(arrays.items())[:20])
        result["sample"] = _sample({k: data[k] for k in list(data)[:10]})
    else:
        result["value"] = _sample(data)
    return result


def _inspect_jsonl(path: Path, options: dict[str, Any]) -> dict[str, Any]:
    encoding = _detect_encoding(_read_head(path, 64 * 1024))
    records: list[Any] = []
    bad = 0
    with open(path, encoding=encoding, errors="replace") as handle:
        for line in handle:
            if len(records) >= 20:
                break
            if not line.strip():
                continue
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError:
                bad += 1
    return {
        "root_type": "json_lines",
        "lines": _count_lines(path),
        "unparsed_sample_lines": bad,
        **_describe_records(records),
    }


def _inspect_text(path: Path, options: dict[str, Any]) -> dict[str, Any]:
    raw = _read_head(path)
    encoding = _detect_encoding(raw)
    text = raw.decode(encoding, errors="replace")
    return {
        "encoding": encoding,
        "lines": _count_lines(path),
        "language": _language(text),
        "sample_text": _clip(text),
    }


def _inspect_binary(path: Path, options: dict[str, Any]) -> dict[str, Any]:
    return {
        "first_bytes_hex": _read_head(path, 32).hex(),
        "suggestion": "Not a document type the inspector reads; use run_python to parse it.",
    }


_INSPECTORS = {
    "pdf": _inspect_pdf,
    "xlsx": _inspect_xlsx,
    "xls": _inspect_xls,
    "csv": _inspect_csv,
    "docx": _inspect_docx,
    "json": _inspect_json,
    "jsonl": _inspect_jsonl,
    "text": _inspect_text,
    "zip": _inspect_binary,
    "binary": _inspect_binary,
}


def inspect(
    path: str,
    name: str | None = None,
    mime: str | None = None,
    options: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Summarize one file: type and metadata, structure, a short sample."""
    file_path = Path(path)
    options = options or {}
    base: dict[str, Any] = {
        "name": name or file_path.name,
        "mime": mime,
        "size_bytes": file_path.stat().st_size,
    }
    try:
        kind = detect_kind(file_path, name, mime)
    except OSError as exc:
        return {**base, "kind": "unknown", "error": f"Could not read the file: {exc}"}
    base["kind"] = kind
    try:
        return {**base, **_INSPECTORS[kind](file_path, options)}
    except Exception as exc:
        return {
            **base,
            "error": f"{type(exc).__name__}: {_clip(str(exc), 500)}",
            "suggestion": "The file may be corrupt or mislabelled. Try run_python to open it "
            "another way, or ask for a fresh copy.",
        }


__all__ = ["inspect", "detect_kind"]

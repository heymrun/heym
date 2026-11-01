"""File sandbox inspector, run in-process against generated fixtures."""

import json
import tempfile
import unittest
import unittest.mock
from pathlib import Path
from types import SimpleNamespace

import docx
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter

from app.services.file_sandbox import inspector
from app.services.file_sandbox.inspector import detect_kind, inspect


def _text_pdf(path: Path, text: str) -> None:
    """Write a one-page PDF whose page has a real text layer."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    path.write_bytes(bytes(out))


class InspectorTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()


class PdfTests(InspectorTestCase):
    def test_text_pdf_reports_pages_text_and_language(self) -> None:
        path = self.dir / "invoice.pdf"
        _text_pdf(path, "This invoice covers the consulting services delivered in March.")
        result = inspect(str(path), name="invoice.pdf")
        self.assertEqual(result["kind"], "pdf")
        self.assertEqual(result["page_count"], 1)
        self.assertTrue(result["text_extractable"])
        self.assertIn("consulting", result["pages"][0]["text"])
        self.assertEqual(result["page_sizes_pt"], [[612.0, 792.0]])
        self.assertEqual(result["language"]["code"], "en")

    def _scan(self, pages: int = 1) -> Path:
        writer = PdfWriter()
        for _ in range(pages):
            writer.add_blank_page(width=595, height=842)
        path = self.dir / "scan.pdf"
        with open(path, "wb") as handle:
            writer.write(handle)
        return path

    def _engine(self, failure: Exception | None = None) -> unittest.mock.Mock:
        def pdf_to_text(_data: bytes, *, language: str, page_range: str) -> SimpleNamespace:
            if failure is not None:
                raise failure
            first, last = (int(n) for n in page_range.split("-"))
            pages = [
                SimpleNamespace(page=n, text=f"GENEL TOPLAM 18.816,00 TL sayfa {n}")
                for n in range(first, last + 1)
            ]
            return SimpleNamespace(language="Latin", pages=pages)

        return unittest.mock.Mock(
            AUTO_LANGUAGE="auto", pdf_to_text=unittest.mock.Mock(wraps=pdf_to_text)
        )

    def test_a_scanned_pdf_is_read_with_ocr(self) -> None:
        engine = self._engine()
        with unittest.mock.patch.object(inspector, "_ocr_engine", return_value=engine):
            result = inspect(str(self._scan(3)))
        self.assertFalse(result["text_extractable"])
        self.assertTrue(result["scanned"])
        self.assertEqual([p["page"] for p in result["pages"]], [1, 2])
        self.assertTrue(all(p["ocr"] for p in result["pages"]))
        self.assertIn("18.816,00", result["pages"][0]["text"])
        self.assertEqual(result["ocr_language"], "Latin")
        self.assertIn("OCR", result["suggestion"])
        engine.pdf_to_text.assert_called_once()
        self.assertEqual(engine.pdf_to_text.call_args.kwargs["language"], "auto")

    def test_ocr_reads_requested_pages_in_runs_and_keeps_the_language(self) -> None:
        engine = self._engine()
        with unittest.mock.patch.object(inspector, "_ocr_engine", return_value=engine):
            result = inspect(str(self._scan(9)), options={"pages": "1-2,5,7-9"})
        self.assertEqual([p["page"] for p in result["pages"]], [1, 2, 5, 7, 8])
        ranges = [c.kwargs["page_range"] for c in engine.pdf_to_text.call_args_list]
        self.assertEqual(ranges, ["1-2", "5-5", "7-8"])
        languages = [c.kwargs["language"] for c in engine.pdf_to_text.call_args_list]
        self.assertEqual(languages, ["auto", "Latin", "Latin"])

    def test_ocr_failure_points_to_the_converter(self) -> None:
        engine = self._engine(ValueError("OCR: 'tesseract' was not found on PATH"))
        with unittest.mock.patch.object(inspector, "_ocr_engine", return_value=engine):
            result = inspect(str(self._scan()))
        self.assertTrue(result["scanned"])
        self.assertIn("tesseract", result["ocr_error"])
        self.assertIn("pdfToText", result["suggestion"])

    def test_password_protected_pdf_says_so(self) -> None:
        source = self.dir / "plain.pdf"
        _text_pdf(source, "Confidential salary table")
        writer = PdfWriter(clone_from=PdfReader(str(source)))
        writer.encrypt("secret")
        path = self.dir / "locked.pdf"
        with open(path, "wb") as handle:
            writer.write(handle)
        result = inspect(str(path))
        self.assertTrue(result["encrypted"])
        self.assertIn("password", result["error"])

    def test_page_range_is_capped(self) -> None:
        path = self.dir / "invoice.pdf"
        _text_pdf(path, "Single page document with enough text to extract.")
        result = inspect(str(path), options={"pages": "1-50"})
        self.assertEqual([p["page"] for p in result["pages"]], [1])


class SpreadsheetTests(InspectorTestCase):
    def _workbook(self) -> Path:
        workbook = Workbook()
        sales = workbook.active
        sales.title = "Sales"
        sales.append(["Region", "Amount", "Closed"])
        for index in range(12):
            sales.append([f"R{index}", 100 + index, index % 2 == 0])
        summary = workbook.create_sheet("Summary")
        summary.append(["Total"])
        summary.append([1266])
        path = self.dir / "q3.xlsx"
        workbook.save(path)
        return path

    def test_lists_sheets_columns_types_and_samples(self) -> None:
        result = inspect(str(self._workbook()))
        self.assertEqual(result["kind"], "xlsx")
        self.assertEqual(result["sheet_names"], ["Sales", "Summary"])
        sales = result["sheets"][0]
        self.assertEqual(sales["row_count"], 13)
        self.assertEqual(
            [(c["name"], c["type"]) for c in sales["columns"]],
            [("Region", "string"), ("Amount", "int"), ("Closed", "bool")],
        )
        self.assertEqual(sales["sample_rows"][0], ["R0", 100, True])
        self.assertEqual(sales["last_filled_row"], 13)

    def test_one_sheet_can_be_requested(self) -> None:
        result = inspect(str(self._workbook()), options={"sheet": "Summary"})
        by_name = {s["name"]: s for s in result["sheets"]}
        self.assertNotIn("columns", by_name["Sales"])
        self.assertEqual(by_name["Sales"]["column_count"], 3)
        self.assertEqual(by_name["Summary"]["sample_rows"], [[1266]])

    def test_unknown_sheet_names_the_real_ones(self) -> None:
        result = inspect(str(self._workbook()), options={"sheet": "Nope"})
        self.assertIn("Sales", result["suggestion"])

    def test_corrupt_workbook_returns_an_error_not_an_exception(self) -> None:
        path = self.dir / "broken.xlsx"
        path.write_bytes(b"PK\x03\x04 not really a zip")
        result = inspect(str(path))
        self.assertIn("error", result)
        self.assertIn("suggestion", result)


class HeaderRowTests(InspectorTestCase):
    """People put titles, notes and blank rows above a table; the header can be on any row."""

    def _sheet(self, title: str, rows: dict[str, object]) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = title
        for ref, value in rows.items():
            sheet[ref] = value
        path = self.dir / f"{title}.xlsx"
        workbook.save(path)
        return path

    def _only_sheet(self, path: Path) -> dict:
        return inspect(str(path))["sheets"][0]

    def test_title_and_blank_rows_above_the_header_are_skipped(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Satışlar"
        sheet["A1"] = "2026 Q3 Satış Kayıtları"
        sheet.merge_cells("A1:C1")
        sheet.append([])
        sheet.append(["Bölge", "Adet", "Tutar (TL)"])
        for index in range(8):
            sheet.append(["Ege" if index % 2 else "Marmara", index + 1, 100.5 * (index + 1)])
        path = self.dir / "sales.xlsx"
        workbook.save(path)

        info = self._only_sheet(path)

        self.assertEqual(info["header_row"], 3)
        self.assertEqual(info["title_rows"], [{"row": 1, "text": "2026 Q3 Satış Kayıtları"}])
        self.assertEqual(
            [(c["name"], c["type"]) for c in info["columns"]],
            [("Bölge", "string"), ("Adet", "int"), ("Tutar (TL)", "float")],
        )
        self.assertEqual(info["sample_rows"][0], ["Marmara", 1, 100.5])
        self.assertEqual(info["read_hint"], "pd.read_excel(path, sheet_name='Satışlar', header=2)")

    def test_a_header_far_down_and_to_the_right_with_year_columns(self) -> None:
        path = self._sheet(
            "Rapor",
            {
                "C2": "Bölge Raporu",
                "C3": "Hazırlayan: Finans",
                "D6": "Ciro",
                "C7": "Bölge",
                "D7": 2025,
                "E7": 2026,
                "C8": "Ege",
                "D8": 10.5,
                "E8": 12.25,
                "C9": "Marmara",
                "D9": 20.0,
                "E9": 22.5,
            },
        )

        info = self._only_sheet(path)

        self.assertEqual(info["header_row"], 7)
        self.assertEqual(info["first_column"], "C")
        self.assertEqual([c["name"] for c in info["columns"]], ["Bölge", 2025, 2026])
        self.assertEqual(
            [t["text"] for t in info["title_rows"]], ["Bölge Raporu", "Hazırlayan: Finans", "Ciro"]
        )
        self.assertEqual(
            info["read_hint"], "pd.read_excel(path, sheet_name='Rapor', header=6, usecols=\"C:E\")"
        )

    def test_the_read_hint_reads_the_table_with_pandas(self) -> None:
        import pandas as pd

        path = self._sheet(
            "Data",
            {"B1": "Title", "B4": "Name", "C4": "Score", "B5": "a", "C5": 1, "B6": "b", "C6": 2},
        )
        info = self._only_sheet(path)
        frame = pd.read_excel(
            path,
            sheet_name="Data",
            header=info["header_row"] - 1,
            usecols=f"{info['first_column']}:C",
        )
        self.assertEqual(list(frame.columns), ["Name", "Score"])
        self.assertEqual(frame["Score"].tolist(), [1, 2])

    def test_a_summary_sheet_returns_label_value_cells(self) -> None:
        path = self._sheet(
            "Özet",
            {
                "A1": "Q3 2026 Özet",
                "A3": "Toplam Ciro",
                "B3": 24851025,
                "A4": "Sipariş Sayısı",
                "B4": 4800,
                "A6": "Hazırlayan",
                "B6": "Finans Ekibi",
            },
        )

        info = self._only_sheet(path)

        self.assertEqual(info["layout"], "key_value")
        self.assertIsNone(info["header_row"])
        self.assertEqual(info["title_rows"], [{"row": 1, "text": "Q3 2026 Özet"}])
        self.assertEqual(
            info["key_values"],
            [
                {"cell": "B3", "label": "Toplam Ciro", "value": 24851025},
                {"cell": "B4", "label": "Sipariş Sayısı", "value": 4800},
                {"cell": "B6", "label": "Hazırlayan", "value": "Finans Ekibi"},
            ],
        )

    def test_label_value_pairs_with_text_first_are_not_a_header(self) -> None:
        path = self._sheet(
            "Pairs",
            {"A1": "Dönem", "B1": "Q3", "A2": "Bölge", "B2": "Tümü", "A3": "Toplam", "B3": 1000},
        )
        info = self._only_sheet(path)
        self.assertEqual(info["layout"], "key_value")
        self.assertEqual(info["key_values"][2], {"cell": "B3", "label": "Toplam", "value": 1000})

    def test_a_real_two_column_table_keeps_its_header(self) -> None:
        path = self._sheet(
            "Two", {"A1": "Bölge", "B1": "Ciro", "A2": "Ege", "B2": 10, "A3": "Marmara", "B3": 20}
        )
        info = self._only_sheet(path)
        self.assertEqual(info["header_row"], 1)
        self.assertEqual([c["name"] for c in info["columns"]], ["Bölge", "Ciro"])
        self.assertNotIn("layout", info)

    def test_an_all_text_table_with_the_header_on_row_one(self) -> None:
        path = self._sheet(
            "People",
            {
                "A1": "Ad",
                "B1": "Şehir",
                "C1": "E-posta",
                "A2": "Ayşe",
                "B2": "İzmir",
                "C2": "a@x.com",
                "A3": "Can",
                "B3": "Ankara",
                "C3": "c@x.com",
            },
        )
        info = self._only_sheet(path)
        self.assertEqual(info["header_row"], 1)
        self.assertEqual(info["title_rows"], [])


class CsvTests(InspectorTestCase):
    def test_turkish_semicolon_csv_in_cp1254(self) -> None:
        path = self.dir / "musteriler.csv"
        rows = ["Ad;Şehir;Tutar"] + [f"Müşteri {i};İstanbul;{i * 10}" for i in range(30)]
        path.write_bytes("\n".join(rows).encode("cp1254"))
        result = inspect(str(path))
        self.assertEqual(result["kind"], "csv")
        self.assertEqual(result["encoding"], "cp1254")
        self.assertEqual(result["delimiter"], ";")
        self.assertEqual(result["rows"], 30)
        self.assertEqual([c["name"] for c in result["columns"]], ["Ad", "Şehir", "Tutar"])
        self.assertEqual(result["sample_rows"][0][1], "İstanbul")

    def test_full_scan_counts_rows_and_nulls_exactly(self) -> None:
        path = self.dir / "data.csv"
        path.write_text("a,b\n1,\n2,3\n", encoding="utf-8")
        result = inspect(str(path), options={"full": True})
        self.assertTrue(result["rows_exact"])
        self.assertEqual(result["rows"], 2)
        self.assertEqual(result["columns"][1]["null_ratio"], 0.5)


class DocxTests(InspectorTestCase):
    def test_outline_tables_and_counts(self) -> None:
        document = docx.Document()
        document.core_properties.author = "Ada"
        document.add_heading("Service Agreement", level=1)
        document.add_paragraph("The supplier delivers the services described below every month.")
        document.add_heading("Payment terms", level=2)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Item"
        table.cell(0, 1).text = "Price"
        path = self.dir / "contract.docx"
        document.save(path)
        result = inspect(str(path))
        self.assertEqual(result["kind"], "docx")
        self.assertEqual(
            result["headings"],
            [{"level": 1, "text": "Service Agreement"}, {"level": 2, "text": "Payment terms"}],
        )
        self.assertEqual(result["tables"], [{"rows": 2, "columns": 2}])
        self.assertEqual(result["first_table_rows"][0], ["Item", "Price"])
        self.assertEqual(result["metadata"]["author"], "Ada")
        self.assertGreater(result["words"], 10)


class JsonTests(InspectorTestCase):
    def test_array_of_records(self) -> None:
        path = self.dir / "orders.json"
        path.write_text(json.dumps([{"id": i, "total": i * 1.5} for i in range(1000)]))
        result = inspect(str(path))
        self.assertEqual(result["root_type"], "array")
        self.assertEqual(result["length"], 1000)
        self.assertEqual(result["record_keys"], {"id": "number", "total": "number"})
        self.assertEqual(len(result["sample_records"]), 3)

    def test_object_root_lists_keys_and_array_lengths(self) -> None:
        path = self.dir / "export.json"
        path.write_text(json.dumps({"meta": {"v": 1}, "items": [1, 2, 3]}))
        result = inspect(str(path))
        self.assertEqual(result["keys"], {"meta": "object", "items": "array"})
        self.assertEqual(result["array_lengths"], {"items": 3})

    def test_large_array_reads_only_the_head(self) -> None:
        path = self.dir / "big.json"
        path.write_text(json.dumps([{"n": i} for i in range(50)]))
        with unittest.mock.patch("app.services.file_sandbox.inspector.JSON_FULL_LOAD_BYTES", 10):
            result = inspect(str(path))
        self.assertIsNone(result["length"])
        self.assertEqual(result["sample_records"][0], {"n": 0})

    def test_json_lines(self) -> None:
        path = self.dir / "events.jsonl"
        path.write_text("\n".join(json.dumps({"event": i}) for i in range(5)) + "\n")
        result = inspect(str(path))
        self.assertEqual(result["root_type"], "json_lines")
        self.assertEqual(result["lines"], 5)


class DetectionTests(InspectorTestCase):
    def test_magic_bytes_win_without_an_extension(self) -> None:
        path = self.dir / "upload"
        _text_pdf(path, "Detected by magic bytes")
        self.assertEqual(detect_kind(path), "pdf")

    def test_binary_gets_a_hint(self) -> None:
        path = self.dir / "blob.bin"
        path.write_bytes(bytes(range(256)))
        result = inspect(str(path))
        self.assertEqual(result["kind"], "binary")
        self.assertIn("run_python", result["suggestion"])


if __name__ == "__main__":
    unittest.main()
